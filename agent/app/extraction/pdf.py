from pathlib import Path

import pdfplumber
import pytesseract

from PIL import Image
from pdf2image import convert_from_path
from docx import Document


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".png",
    ".jpg",
    ".jpeg",
    ".tiff",
    ".bmp",
}


def extract_document(file_path: Path) -> tuple[str, str]:
    """
    Extract complete raw text from a supported document.

    Returns:
        text: Extracted document text.
        extraction_method: Method used.
    """

    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension}"
        )

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    return extract_image(file_path)


def extract_docx(file_path: Path) -> tuple[str, str]:
    """
    Extract text from a DOCX.

    Strategy:
    1. Extract paragraphs and tables.
    2. If meaningful text is found, return it.
    3. Otherwise OCR embedded images.
    """

    document = Document(file_path)

    sections = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            sections.append(text)

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            row_text = " | ".join(
                cell for cell in cells if cell
            )

            if row_text:
                sections.append(row_text)

    direct_text = "\n".join(sections).strip()

    # DOCX contains usable text
    if len(direct_text) >= 20:
        return direct_text, "docx"

    # No meaningful text → OCR embedded images
    return extract_docx_with_ocr(document)

def extract_docx_with_ocr(document: Document) -> tuple[str, str]:
    """
    OCR images embedded inside a DOCX document.
    """

    extracted_images = []

    for relationship in document.part.rels.values():

        if "image" not in relationship.target_ref:
            continue

        image_data = relationship.target_part.blob

        from io import BytesIO

        image = Image.open(
            BytesIO(image_data)
        )

        text = pytesseract.image_to_string(image).strip()

        if text:
            extracted_images.append(text)

    extracted_text = "\n\n".join(
        extracted_images
    ).strip()

    return extracted_text, "ocr"


def extract_pdf(file_path: Path) -> tuple[str, str]:
    """
    Try direct PDF text extraction first.

    If insufficient text is extracted, fall back to OCR.
    """

    extracted_pages = []

    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()

            if text:
                extracted_pages.append(text)

    direct_text = "\n\n".join(extracted_pages).strip()

    # Direct extraction succeeded.
    if direct_text:
        return direct_text, "pdf"

    # No useful text → OCR fallback.
    return extract_pdf_with_ocr(file_path)


def extract_pdf_with_ocr(file_path: Path) -> tuple[str, str]:
    """
    Convert PDF pages to images and run Tesseract OCR.
    """

    pages = convert_from_path(
        file_path,
        dpi=300,
    )

    extracted_pages = []

    for page_number, page_image in enumerate(pages, start=1):
        text = pytesseract.image_to_string(
            page_image
        )

        if text:
            extracted_pages.append(
                f"--- Page {page_number} ---\n{text}"
            )

    extracted_text = "\n\n".join(
        extracted_pages
    ).strip()

    return extracted_text, "ocr"


def extract_image(file_path: Path) -> tuple[str, str]:
    """
    Run Tesseract OCR on an image.
    """

    image = Image.open(file_path)

    text = pytesseract.image_to_string(
        image
    )

    return text.strip(), "ocr"